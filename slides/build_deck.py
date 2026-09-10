#!/usr/bin/env python3
"""Build a clean teaching deck without modifying the supplied W7.pptx."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
NAVY = RGBColor(31, 78, 121)
BLUE = RGBColor(46, 117, 182)
TEXT = RGBColor(45, 45, 45)
MUTED = RGBColor(110, 110, 110)
LIGHT = RGBColor(235, 243, 250)
YELLOW = RGBColor(255, 242, 204)
RED = RGBColor(180, 40, 40)
WHITE = RGBColor(255, 255, 255)


def textbox(slide, text, x, y, w, h, *, size=20, color=TEXT, bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.03)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def box(slide, x, y, w, h, fill=LIGHT, line=BLUE, radius=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def title(slide, text):
    textbox(slide, text, 0.55, 0.25, 8.9, 0.65, size=26, color=NAVY, bold=True)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.02), Inches(8.9), Inches(0.02))
    line.fill.solid(); line.fill.fore_color.rgb = RGBColor(204, 204, 204); line.line.fill.background()


def bullets(slide, items, x=0.65, y=1.3, w=8.6, h=3.8, size=20):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame; tf.clear(); tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item; p.level = 0; p.font.name = "Arial"; p.font.size = Pt(size); p.font.color.rgb = TEXT
        p.space_after = Pt(10); p._p.get_or_add_pPr().insert(0, p._p._new_buChar()) if False else None
        p.text = "• " + item
    return shape


def table(slide, rows, x, y, w, h, font=14):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = shape.table
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = tbl.cell(r, c); cell.text = str(value)
            cell.margin_left = Inches(0.05); cell.margin_right = Inches(0.05)
            cell.fill.solid(); cell.fill.fore_color.rgb = NAVY if r == 0 else WHITE
            for p in cell.text_frame.paragraphs:
                p.font.name = "Arial"; p.font.size = Pt(font); p.font.color.rgb = WHITE if r == 0 else TEXT; p.font.bold = r == 0
    return shape


def build(result_path: Path | None = None) -> Path:
    prs = Presentation(); prs.slide_width = Inches(10); prs.slide_height = Inches(5.625)
    blank = prs.slide_layouts[6]
    # 1
    s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = NAVY
    textbox(s, "Three product APIs become one\nreproducible K-means lesson", 0.7, 1.2, 8.5, 1.4, size=32, color=WHITE, bold=True)
    textbox(s, "Internet Programming · price + inventory clustering · 2026", 0.72, 3.0, 8.5, 0.4, size=18, color=RGBColor(200, 220, 240))
    textbox(s, "Three cloud APIs → validation → StandardScaler → K-means → report", 0.72, 3.65, 8.5, 0.7, size=20, color=WHITE)
    # 2
    s = prs.slides.add_slide(blank); title(s, "The lesson asks whether API records can be clustered by price and inventory")
    bullets(s, ["Collect products from one existing TA API and two simulated-group APIs running on the cloud.", "Every source provides price and stock directly through HTTP.", "Use machine learning to learn memberships; grouping is not a query or category rule.", "Explain what the clusters can and cannot mean with only a small snapshot."])
    # 3
    s = prs.slides.add_slide(blank); title(s, "Three public cloud endpoints feed one reproducible local pipeline")
    for index, (x, label, sub) in enumerate([(0.7, "std6730202734 (TA)", "Cloud HTTP :3067"), (0.7, "demo-shop-b", "Cloud HTTP :3120 · synthetic data"), (0.7, "demo-shop-c", "Cloud HTTP :3121 · synthetic data")]):
        yy = 1.35 + index * 1.05
        box(s, x, yy, 2.65, 0.75, LIGHT, BLUE, True); textbox(s, label, x + 0.12, yy + 0.12, 2.4, 0.25, size=17, color=NAVY, bold=True); textbox(s, sub, x + 0.12, yy + 0.42, 2.4, 0.2, size=12, color=MUTED)
        textbox(s, "→", 3.55, yy + 0.2, 0.45, 0.35, size=28, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    box(s, 4.1, 1.75, 2.35, 2.0, LIGHT, BLUE, True); textbox(s, "Local\naggregator", 4.25, 2.1, 2.05, 0.6, size=23, color=NAVY, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE); textbox(s, "fetch · validate\nnormalize", 4.3, 2.85, 1.95, 0.5, size=16, color=TEXT, align=PP_ALIGN.CENTER)
    textbox(s, "→", 6.7, 2.45, 0.45, 0.35, size=28, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    box(s, 7.2, 1.75, 2.1, 2.0, YELLOW, BLUE, True); textbox(s, "Snapshot\n+ K-means", 7.35, 2.15, 1.8, 0.65, size=22, color=NAVY, bold=True, align=PP_ALIGN.CENTER); textbox(s, "CSV · JSON · plots", 7.4, 3.0, 1.7, 0.3, size=15, color=TEXT, align=PP_ALIGN.CENTER)
    textbox(s, "The teaching run needs no root login or SSH tunnel; it performs ordinary HTTP GET requests.", 0.7, 4.75, 8.7, 0.35, size=15, color=MUTED)
    # 4
    s = prs.slides.add_slide(blank); title(s, "A shared API shape makes aggregation clear while preserving source identity")
    table(s, [["Source", "Response shape", "Rows", "Data origin"], ["std6730202734", "{data, meta}", "4", "existing TA API"], ["demo-shop-b", "{source_id, data, meta}", "4", "synthetic classroom"], ["demo-shop-c", "{source_id, data, meta}", "4", "synthetic classroom"]], 0.65, 1.35, 8.7, 2.2, 15)
    box(s, 0.8, 3.95, 8.3, 0.85, LIGHT, BLUE, True); textbox(s, "Canonical key = source_id + ':' + source_product_id", 1.0, 4.12, 7.9, 0.28, size=22, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    textbox(s, "A product with the same local ID in two stores stays as two records.", 0.85, 5.0, 8.2, 0.3, size=17, color=TEXT, align=PP_ALIGN.CENTER)
    # 5
    s = prs.slides.add_slide(blank); title(s, "Every training value comes from an API response with visible provenance")
    table(s, [["Field", "Used by model", "Rule"], ["price", "yes", "THB selling price per unit"], ["stock", "yes", "current inventory units"], ["source_id", "no", "provenance and plot marker only"], ["category", "no", "display only"], ["price_is_synthetic", "no", "honest classroom-data label"]], 0.6, 1.3, 8.8, 2.75, 14)
    box(s, 1.0, 4.35, 7.95, 0.75, YELLOW, BLUE, True); textbox(s, "No central price file: all 12 [price, stock] vectors arrive directly from HTTP APIs.", 1.15, 4.55, 7.65, 0.3, size=17, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    textbox(s, "The two simulated catalogs are labelled synthetic; the services and HTTP calls are live.", 1.0, 5.15, 7.95, 0.28, size=15, color=RED, bold=True, align=PP_ALIGN.CENTER)
    # 6
    s = prs.slides.add_slide(blank); title(s, "Each accepted product becomes one two-dimensional vector after validation")
    box(s, 0.8, 1.45, 3.0, 2.35, LIGHT, BLUE, True); textbox(s, "Product record", 1.0, 1.7, 2.6, 0.3, size=21, color=NAVY, bold=True, align=PP_ALIGN.CENTER); textbox(s, "price = 1,590 THB\nstock = 42 units", 1.1, 2.35, 2.4, 0.65, size=25, color=TEXT, align=PP_ALIGN.CENTER)
    textbox(s, "→", 4.05, 2.25, 0.55, 0.4, size=30, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    box(s, 4.75, 1.45, 4.3, 2.35, YELLOW, BLUE, True); textbox(s, "xᵢ = [price, stock]", 5.0, 1.95, 3.8, 0.45, size=28, color=NAVY, bold=True, align=PP_ALIGN.CENTER); textbox(s, "StandardScaler fits once\non the combined dataset", 5.15, 2.75, 3.5, 0.55, size=19, color=TEXT, align=PP_ALIGN.CENTER)
    bullets(s, ["Source, category, product ID and mock flag are display/provenance fields.", "Stock is inventory, never monthly sales.", "Scaling prevents price units from dominating distance."], x=0.9, y=4.25, w=8.2, h=1.0, size=16)
    # 7
    s = prs.slides.add_slide(blank); title(s, "K-means learns memberships and k is evaluated on the same snapshot")
    bullets(s, ["Fit StandardScaler once, then KMeans(random_state=42, n_init=10).", "Use fit_predict for labels; groupby only summarizes learned labels afterward.", "Compare k = 2…min(6, n−1, unique vectors) with inertia, silhouette, and membership inspection.", "Cluster IDs are arbitrary labels, not ranks or business recommendations."], x=0.7, y=1.35, w=8.7, h=2.7, size=18)
    box(s, 1.0, 4.25, 7.9, 0.65, LIGHT, BLUE, True); textbox(s, "First teaching run: k=3 is pedagogical; three APIs do not determine k=3.", 1.15, 4.45, 7.6, 0.25, size=18, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    # 8: actual demo output, only when the caller explicitly supplies a recorded run.
    if result_path and result_path.exists():
        result = json.loads(result_path.read_text(encoding="utf-8"))
        scatter = result_path.parent / "price-stock-clusters.png"
        if scatter.exists():
            s = prs.slides.add_slide(blank); title(s, "The live cloud run produces learned memberships and centroids")
            s.shapes.add_picture(str(scatter), Inches(0.45), Inches(1.25), width=Inches(6.15))
            box(s, 6.85, 1.45, 2.55, 2.75, YELLOW, BLUE, True)
            textbox(s, "Demo result", 7.05, 1.7, 2.15, 0.3, size=20, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
            textbox(s, f"Mode\n{result.get('mode', 'DEMO')}\n\nMatrix\n{result['matrix_shape'][0]} × {result['matrix_shape'][1]}\n\nk = {result['selected']['k']}\nSilhouette = {result['selected']['silhouette']:.3f}", 7.1, 2.15, 2.05, 1.75, size=16, color=TEXT, align=PP_ALIGN.CENTER)
            textbox(s, "B/C catalogs are synthetic; HTTP and K-means are live.", 6.9, 4.5, 2.4, 0.5, size=14, color=RED, bold=True, align=PP_ALIGN.CENTER)
    # 8
    s = prs.slides.add_slide(blank); title(s, "The current live run is complete and k=3 is supported by this snapshot")
    table(s, [["Evidence", "Observed"], ["Cloud sources", "3 HTTP endpoints"], ["Products", "12 (4 + 4 + 4)"], ["Feature matrix", "12 × 2 [price, stock]"], ["Rejected rows", "0"], ["k=3 silhouette", "0.587 (highest candidate now)"]], 0.75, 1.35, 8.45, 2.65, 16)
    box(s, 0.95, 4.35, 8.1, 0.72, YELLOW, BLUE, True); textbox(s, "Cluster sizes: 3 / 4 / 5 — members cross API boundaries.", 1.1, 4.56, 7.8, 0.25, size=21, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    textbox(s, "If API data changes, diagnostics and memberships can change; every run keeps its own snapshot.", 0.9, 5.15, 8.2, 0.28, size=15, color=MUTED, align=PP_ALIGN.CENTER)
    # 9
    s = prs.slides.add_slide(blank); title(s, "The live demonstration follows one visible path from JSON to learned labels")
    table(s, [["Step", "Action", "Badge"], ["1", "Open :3067, :3120 and :3121; show 4 rows each", "LIVE"], ["2", "Point to price, stock and synthetic provenance", "LIVE"], ["3", "Run ./scripts/run_demo.sh --mode live --k 3", "LIVE"], ["4", "Open report, centroids and candidate-k diagnostics", "LIVE"], ["5", "Use archived capture only if network is unavailable", "SNAPSHOT"]], 0.5, 1.25, 9.0, 3.1, 13)
    textbox(s, "No SSH tunnel, database edit or manual category grouping is needed during class.", 0.8, 4.8, 8.4, 0.35, size=19, color=RED, bold=True, align=PP_ALIGN.CENTER)
    # 10
    s = prs.slides.add_slide(blank); title(s, "The small dataset supports a teaching demonstration, not production decisions")
    bullets(s, ["Twelve rows are enough to show vectors, scaling, labels, centroids, and k diagnostics.", "A cluster does not prove popularity, turnover, urgency, profit, or a reorder policy.", "Eight demo B/C rows are visibly synthetic although their cloud APIs are live.", "A future text mode should be separate from price/inventory clustering and use documented TF-IDF or embeddings."], x=0.7, y=1.3, w=8.7, h=3.1, size=19)
    # 11
    s = prs.slides.add_slide(blank); title(s, "The handoff is runnable now with a tested live path and snapshot fallback")
    bullets(s, ["Check the three browser URLs; each must return four price/stock records.", "Run tests: .venv/bin/python -m unittest discover -s tests -v", "Run ./scripts/run_demo.sh --mode live --k 3; use snapshot mode only as fallback.", "Review outputs/<run_id>/report.html and speaker notes before teaching."], x=0.7, y=1.25, w=8.7, h=2.5, size=19)
    textbox(s, "References", 0.7, 4.0, 2.0, 0.3, size=20, color=BLUE, bold=True)
    textbox(s, "scikit-learn KMeans · StandardScaler · silhouette_score\nhttps://scikit-learn.org/stable/\nSource data: three cloud HTTP endpoints; demo B/C catalogs are synthetic", 0.7, 4.35, 8.3, 0.8, size=15, color=MUTED)
    out = ROOT / "slides" / "W7_API_KMeans_Revised.pptx"
    prs.save(out)
    return out


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--result", type=Path, help="recorded clustering.result.json to embed")
    args = parser.parse_args()
    build(args.result)
